import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    COLOR_BG = RGBColor(248, 250, 252)       # Light Slate
    COLOR_PRIMARY = RGBColor(15, 23, 42)     # Deep Navy
    COLOR_ACCENT = RGBColor(37, 99, 235)     # Royal Blue
    COLOR_TEAL = RGBColor(20, 184, 166)      # Emerald/Teal
    COLOR_MUTED = RGBColor(100, 116, 139)    # Slate Muted
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_CARD_BORDER = RGBColor(226, 232, 240)

    base_dir = os.path.abspath(os.path.dirname(__file__))
    img_dir = os.path.join(base_dir, "presentation_screenshots")

    def add_header(slide, title, category=None):
        # Header category tag
        if category:
            cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
            tf = cat_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = category.upper()
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = COLOR_ACCENT

        # Title
        top_pos = Inches(0.75) if category else Inches(0.5)
        title_box = slide.shapes.add_textbox(Inches(0.8), top_pos, Inches(11.7), Inches(0.6))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(24)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_PRIMARY

        # Thin accent bar under title
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), top_pos + Inches(0.65), Inches(1.5), Inches(0.05))
        bar.fill.solid()
        bar.fill.fore_color.rgb = COLOR_ACCENT
        bar.line.fill.background()

    def add_footer(slide):
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.7), Inches(0.35))
        tf = footer_box.text_frame
        p = tf.paragraphs[0]
        p.text = "SKIT, JAIPUR  |  www.skit.ac.in  |  EXIM 2D World Traders HR ERP Platform"
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_MUTED

    # ==========================================
    # SLIDE 1: TITLE SLIDE
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    # Background accent card
    bg = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_PRIMARY
    bg.line.fill.background()

    # Title card box
    tb = s1.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.333), Inches(2.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    p0.text = "SWAMI KESHVANAND INSTITUTE OF TECHNOLOGY, MANAGEMENT & GRAMOTHAN, JAIPUR"
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_TEAL

    p1 = tf.add_paragraph()
    p1.text = "Industrial Training & Major Project Presentation"
    p1.font.size = Pt(13)
    p1.font.color.rgb = RGBColor(148, 163, 184)
    p1.space_before = Pt(8)

    p2 = tf.add_paragraph()
    p2.text = "EXIM 2D World Traders\nCloud-Native HR ERP & In-App AI Copilot Platform"
    p2.font.size = Pt(30)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_WHITE
    p2.space_before = Pt(16)

    # Submitted By & To cards
    c1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(4.5), Inches(5.3), Inches(2.2))
    c1.fill.solid()
    c1.fill.fore_color.rgb = RGBColor(30, 41, 59)
    c1.line.color.rgb = RGBColor(51, 65, 85)
    tf1 = c1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "SUBMITTED BY:"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEAL
    
    p = tf1.add_paragraph()
    p.text = "• Presenter 1 (Frontend & UI/UX Architecture)\n• Presenter 2 (Backend, Database Systems & AI)"
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_WHITE
    p.space_before = Pt(6)

    p = tf1.add_paragraph()
    p.text = "Department of Computer Science & Engineering"
    p.font.size = Pt(11)
    p.font.color.rgb = RGBColor(148, 163, 184)
    p.space_before = Pt(6)

    c2 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.0), Inches(4.5), Inches(5.3), Inches(2.2))
    c2.fill.solid()
    c2.fill.fore_color.rgb = RGBColor(30, 41, 59)
    c2.line.color.rgb = RGBColor(51, 65, 85)
    tf2 = c2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "SUBMITTED TO:"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT
    
    p = tf2.add_paragraph()
    p.text = "Dr. Meenakshi Nawal"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE
    p.space_before = Pt(6)

    p = tf2.add_paragraph()
    p.text = "Associate Professor\nDepartment of Computer Science & Engineering\nSKIT, Jaipur"
    p.font.size = Pt(12)
    p.font.color.rgb = RGBColor(148, 163, 184)
    p.space_before = Pt(4)

    # ==========================================
    # SLIDE 2: INDEX
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Table of Contents", "Overview")
    add_footer(s2)

    topics = [
        ("01", "Company & Business Domain", "EXIM 2D Global Trade & Logistics"),
        ("02", "Technology Stack Used", "React 19, Node.js, Express, MongoDB, EXIM-AI"),
        ("03", "Goal & Objectives of Training", "Enterprise Full-Stack & Dual-Portal Architecture"),
        ("04", "System Architecture & Pipeline", "Decoupled Client-Server & Security Guards"),
        ("05", "Sample Database Ingestion", "Enterprise Indian Workforce Dataset (Kaggle)"),
        ("06", "Executive HR Command Center", "Workforce Telemetry & Global Hub Radar"),
        ("07", "Attendance & Workforce Modules", "Biometric Logs, Directory & Onboarding"),
        ("08", "Employee Self-Service (ESS)", "Clock-In, Quota Balances & INR Payslips"),
        ("09", "EXIM-AI Contextual Copilot", "Role-Aware Assistant & Action Dispatcher"),
        ("10", "Future Scope & Conclusion", "IoT Geofencing, Banking APIs & Q&A")
    ]

    for i, (num, title, desc) in enumerate(topics):
        col = i // 5
        row = i % 5
        left = Inches(0.8 + col * 5.9)
        top = Inches(1.8 + row * 0.95)

        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.6), Inches(0.8))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_WHITE
        card.line.color.rgb = COLOR_CARD_BORDER
        tf = card.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = f"{num}.  {title}"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = COLOR_PRIMARY

        p2 = tf.add_paragraph()
        p2.text = f"     {desc}"
        p2.font.size = Pt(10)
        p2.font.color.rgb = COLOR_MUTED

    # ==========================================
    # SLIDE 3: COMPANY INTRODUCTION
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Company Introduction & Domain Context", "Organization Background")
    add_footer(s3)

    c = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.8))
    c.fill.solid()
    c.fill.fore_color.rgb = COLOR_WHITE
    c.line.color.rgb = COLOR_CARD_BORDER
    tf = c.text_frame
    tf.word_wrap = True

    pts = [
        ("Company Profile:", "EXIM 2D World Traders is a premier international enterprise specializing in global import-export trade, maritime cargo logistics, container fleet management, and large-scale civil infrastructure projects."),
        ("Key Operating Hubs:", "Major maritime terminals and project divisions including JNPT Port Mumbai, Mundra Port Gujarat, Chennai Port Hub, Kolkata Logistics Hub, and Delhi Corporate Headquarters."),
        ("Workforce Demographics:", "Over 140+ multidisciplinary personnel comprising port logistics managers, customs clearance officers, civil site engineers, heavy equipment supervisors, and finance executives."),
        ("Core Operational Challenge:", "Manual legacy spreadsheets and disparate tools led to shift communication gaps, attendance discrepancies, delayed leave approvals, and sluggish payroll disbursal across international hubs."),
        ("The Strategic Solution:", "Development of a unified, cloud-native HR ERP system featuring dual-persona access (Executive HR & Staff ESS) combined with an integrated conversational AI Copilot (EXIM-AI).")
    ]
    for i, (head, body) in enumerate(pts):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"•  {head} "
        p.font.bold = True
        p.font.size = Pt(14)
        p.font.color.rgb = COLOR_ACCENT
        if i > 0: p.space_before = Pt(12)

        run = p.add_run()
        run.text = body
        run.font.bold = False
        run.font.size = Pt(13)
        run.font.color.rgb = COLOR_PRIMARY

    # ==========================================
    # SLIDE 4: TECHNOLOGY STACK USED
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Technology Stack Used", "System Architecture")
    add_footer(s4)

    techs = [
        ("Frontend Architecture", COLOR_ACCENT, [
            ("React 19", "Modern UI library with concurrent rendering & modular component architecture."),
            ("Vite 8", "Next-gen lightning-fast HMR dev server and optimized production bundler."),
            ("Tailwind CSS v4", "Next-generation utility-first styling with glassmorphic backdrop-blur aesthetics."),
            ("Framer Motion 12", "Physics-based spring animations, ambient glowing orbs & smooth modal transitions."),
            ("Recharts 3", "Interactive SVG data visualizations for telemetry & weekly attendance analytics."),
            ("React Router v7", "Declarative client-side routing with protected route security guards.")
        ]),
        ("Backend & Database Systems", COLOR_TEAL, [
            ("Node.js Runtime", "Scalable asynchronous JavaScript server-side execution environment."),
            ("Express.js API", "RESTful routing controllers handling authentication, payroll, and employees."),
            ("MongoDB (NoSQL)", "Flexible document store housing employee records, shifts, and leaves."),
            ("JWT Stateless Auth", "Signed JSON Web Tokens enforcing Role-Based Access Control (RBAC)."),
            ("Axios HTTP Client", "Promise-based client-to-server API communication with bearer headers."),
            ("Database Seeding", "Automated ingestion pipeline populating verified enterprise HR datasets.")
        ]),
        ("EXIM-AI Copilot Engine", RGBColor(147, 51, 234), [
            ("Role-Aware Intent Classifier", "Contextually distinguishes between HR Executives and Field Staff inquiries."),
            ("Conversational Interface", "Instant parsing of natural language queries on leaves, shifts & ports."),
            ("Deep-Linking Action Dispatcher", "Dynamically injects actionable, one-click navigation CTA buttons in chat."),
            ("In-Memory Grounding", "Pre-loaded with accurate domain data (port operations, leave quotas, payroll)."),
            ("Sub-50ms Response", "High-speed client-side intent execution delivering zero-latency assistance.")
        ])
    ]

    for i, (sec_title, sec_color, items) in enumerate(techs):
        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 3.95), Inches(1.8), Inches(3.8), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_WHITE
        card.line.color.rgb = sec_color
        tf = card.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = sec_title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = sec_color

        for item_name, item_desc in items:
            pi = tf.add_paragraph()
            pi.text = f"• {item_name}: "
            pi.font.size = Pt(11)
            pi.font.bold = True
            pi.font.color.rgb = COLOR_PRIMARY
            pi.space_before = Pt(8)

            run = pi.add_run()
            run.text = item_desc
            run.font.bold = False
            run.font.size = Pt(10)
            run.font.color.rgb = COLOR_MUTED

    # ==========================================
    # SLIDE 5: GOAL OF TRAINING
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Goal & Core Objectives of Training", "Project Objectives")
    add_footer(s5)

    goals = [
        ("Hands-on Full-Stack Web Mastery", "Develop real-world proficiency in engineering production-grade web platforms using modern decoupled architectures (React 19 + Node.js/Express + MongoDB)."),
        ("Dual-Persona Enterprise Workflows", "Architect and isolate secure, role-specific user journeys for Executive HR governance and Employee Self-Service (ESS) within a single Single Page Application."),
        ("Multi-Hub Data Modeling & Ingestion", "Learn end-to-end data pipelines: transforming raw multi-location CSV datasets (JNPT Mumbai, Mundra Gujarat, Delhi HQ) into structured MongoDB schemas."),
        ("Intelligent Conversational Assistance", "Conceptualize, build, and deploy an in-app AI Copilot (EXIM-AI) capable of parsing user queries and deep-linking directly to target application pages."),
        ("Engineering Best Practices & Git Workflow", "Strengthen professional engineering habits—including Git branching, code reviews, version control, and GitHub repository deployment.")
    ]

    for i, (head, desc) in enumerate(goals):
        box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8 + i * 0.95), Inches(11.7), Inches(0.82))
        box.fill.solid()
        box.fill.fore_color.rgb = COLOR_WHITE
        box.line.color.rgb = COLOR_CARD_BORDER
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"0{i+1}.  {head}"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = COLOR_ACCENT

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_PRIMARY
        p2.space_before = Pt(2)

    # ==========================================
    # SLIDE 6: SYSTEM ARCHITECTURE FLOWCHART
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "System Pipeline & Decoupled Architecture", "Technical Blueprint")
    add_footer(s6)

    pipe_steps = [
        ("1. Presentation Layer", "React 19 + Vite", "Role-switched SPA on Port 5173 rendering Executive HR and ESS portals with Tailwind CSS."),
        ("2. Security & Guard", "JWT Route Protection", "Interception of protected routes via useAuth hook verifying stateless bearer token & roles."),
        ("3. REST API Controllers", "Express.js Services", "Asynchronous endpoints handling auth, employee rosters, biometric logs & payroll calculations."),
        ("4. Data Persistence", "MongoDB Collections", "NoSQL document storage housing 140+ ingested records of staff, attendance & leave quotas."),
        ("5. Context AI Copilot", "EXIM-AI Bridge", "In-app natural language engine resolving queries with instant deep-link navigation shortcuts.")
    ]

    for i, (title, tech, desc) in enumerate(pipe_steps):
        b = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 2.38), Inches(2.2), Inches(2.25), Inches(4.2))
        b.fill.solid()
        b.fill.fore_color.rgb = COLOR_WHITE
        b.line.color.rgb = COLOR_ACCENT if i % 2 == 0 else COLOR_TEAL
        tf = b.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = COLOR_PRIMARY

        p2 = tf.add_paragraph()
        p2.text = tech
        p2.font.size = Pt(11)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_ACCENT
        p2.space_before = Pt(6)

        p3 = tf.add_paragraph()
        p3.text = desc
        p3.font.size = Pt(10)
        p3.font.color.rgb = COLOR_MUTED
        p3.space_before = Pt(10)

    # ==========================================
    # HELPER FOR SCREENSHOT SLIDES
    # ==========================================
    def add_image_slide(title, category, img_filename, key_points):
        s = prs.slides.add_slide(blank_layout)
        add_header(s, title, category)
        add_footer(s)

        img_path = os.path.join(img_dir, img_filename)
        if os.path.exists(img_path):
            # Insert image on left
            s.shapes.add_picture(img_path, Inches(0.8), Inches(1.8), Inches(7.5), Inches(4.8))
        else:
            # Fallback placeholder box
            ph = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(7.5), Inches(4.8))
            ph.fill.solid()
            ph.fill.fore_color.rgb = RGBColor(241, 245, 249)
            ph.text_frame.text = f"[Image: {img_filename}]"

        # Information Card on Right
        info_card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.55), Inches(1.8), Inches(3.95), Inches(4.8))
        info_card.fill.solid()
        info_card.fill.fore_color.rgb = COLOR_WHITE
        info_card.line.color.rgb = COLOR_CARD_BORDER
        tf = info_card.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = "KEY MODULE HIGHLIGHTS"
        p0.font.size = Pt(12)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_ACCENT

        for h, d in key_points:
            p = tf.add_paragraph()
            p.text = f"• {h}: "
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = COLOR_PRIMARY
            p.space_before = Pt(10)

            run = p.add_run()
            run.text = d
            run.font.bold = False
            run.font.size = Pt(10)
            run.font.color.rgb = COLOR_MUTED

    # ==========================================
    # SLIDE 7: SAMPLE HR DATASET (EXCEL VIEW)
    # ==========================================
    add_image_slide(
        "Sample HR Dataset Ingestion (Excel View)",
        "Database Layer",
        "10_sample_dataset_excel.png",
        [
            ("Verified Enterprise Source", "Standard enterprise HR & trade logistics dataset sourced from Kaggle repository."),
            ("Indian Workforce Grounding", "140+ structured records featuring Indian personnel (Aarav Sharma, Priya Patel, Amitabh Verma, Neha Gupta)."),
            ("National Trade Hubs", "Personnel mapped across JNPT Mumbai, Mundra Gujarat, Chennai Port, Kolkata Hub & Delhi HQ."),
            ("Automated Database Seeding", "Transformed and ingested into MongoDB Atlas collections via custom seedData.js pipeline."),
            ("INR Payroll & Metric Realism", "Includes gross monthly salaries (INR), leave quotas, shift timings & biometric punctuality rates.")
        ]
    )

    # ==========================================
    # SLIDE 8: DUAL GATEWAY & LOGIN
    # ==========================================
    add_image_slide(
        "Dual Gateway & Role-Based Authentication",
        "Frontend Module",
        "01_landing_gateway.png",
        [
            ("Gateway Selector", "Intuitive home portal allowing personnel to route to Executive HR or Employee Self-Service."),
            ("Modern Glassmorphism", "Built with Tailwind CSS v4 featuring translucent frosted surfaces and ambient orbs."),
            ("Credential Verification", "Client-side validation with real-time error feedback and show/hide password toggles."),
            ("Stateless Session Token", "Upon successful login, a cryptographic JWT token and userRole are initialized in localStorage."),
            ("Smooth Micro-Interactions", "Fluid card hover and click physics powered by Framer Motion 12.")
        ]
    )

    # ==========================================
    # SLIDE 9: EXECUTIVE HR COMMAND CENTER
    # ==========================================
    add_image_slide(
        "Executive HR Command Center & Radar",
        "HR Operations",
        "03_hr_dashboard.png",
        [
            ("Live Workforce Telemetry", "Real-time counters for Total Workforce (142), Present on Site (90.1%), and Leaves (9)."),
            ("National Site Hub Radar", "Monitors live weather, vessel dockings, and biometrics across JNPT, Mundra, Delhi HQ & Expressway."),
            ("Interactive Recharts SVG", "Weekly attendance trends and departmental headcounts visualized with responsive SVG graphs."),
            ("One-Click Quick Actions", "Direct shortcuts to onboard employees, mark biometrics, approve leaves, and disburse payroll."),
            ("Responsive Collapse", "Dynamic sidebar supporting desktop collapse and mobile drawer transformation.")
        ]
    )

    # ==========================================
    # SLIDE 10: WORKFORCE DIRECTORY & FILTERING
    # ==========================================
    add_image_slide(
        "Workforce Directory & Onboarding Management",
        "HR Operations",
        "04_employees_directory.png",
        [
            ("Multi-Criteria Search", "Instant keyword search across employee names, designations, and department tags."),
            ("Department Filtering", "Dynamic filter tabs for Logistics, Civil Infrastructure, Global Trade, IT & Finance."),
            ("Grid & Table View Switcher", "Seamless toggle between visual profile cards and compact enterprise tabular rows."),
            ("Employee Status Indicators", "Visual badge chips displaying Active, On Leave, and Probationary states."),
            ("New Employee Onboarding", "Interactive modal form with real-time validation to register incoming field engineers.")
        ]
    )

    # ==========================================
    # SLIDE 11: BIOMETRIC ATTENDANCE SHEET
    # ==========================================
    add_image_slide(
        "Biometric Attendance & Shift Tracking",
        "Workforce Governance",
        "05_attendance_sheet.png",
        [
            ("Biometric Clock Logs", "Accurate records of morning, evening, and field shifts across maritime terminals."),
            ("Punctuality Monitoring", "Automatic tagging of On-Time and Late check-ins with exact hours logged."),
            ("Field Terminal Sync", "Simulated hardware sync capturing arrivals at JNPT Port, Mundra, and Delhi Expressway."),
            ("Interactive Manual Check-In", "Modal interface enabling supervisors to manually log field staff clock-ins with timestamps."),
            ("Data Export Readiness", "One-click export capability preparing attendance reports for monthly payroll processing.")
        ]
    )

    # ==========================================
    # SLIDE 12: ENTERPRISE PAYROLL & CURRENCY
    # ==========================================
    add_image_slide(
        "Enterprise Payroll & INR Disbursal",
        "Financial Governance",
        "06_payroll_management.png",
        [
            ("Indian Rupee (INR) Base", "Calculated on a total monthly workforce payroll outflow exceeding ₹2.07 Crores."),
            ("Statutory Tax Compliance", "Automated computation of EPF, ESI, professional tax, and standard deductions."),
            ("One-Click Batch Disbursal", "Simulated bulk bank transfer triggering instantaneous salary credits to 142 personnel."),
            ("Global Currency Switcher", "Built-in multi-currency conversion supporting INR, USD, EUR, SGD, and AED."),
            ("Expense Claims Tracker", "Live monitoring of pending travel and site reimbursement claims awaiting HR audit.")
        ]
    )

    # ==========================================
    # SLIDE 13: EMPLOYEE SELF-SERVICE PORTAL
    # ==========================================
    add_image_slide(
        "Employee Self-Service (ESS) Portal",
        "Staff Workflows",
        "07_employee_dashboard.png",
        [
            ("Personalized Dashboard", "Greets authenticated staff (e.g., Aarav Sharma - EX-101) with their role and unit overview."),
            ("One-Click Shift Clock-In", "Enables workers to record daily shift start and end times with automated clock stamps."),
            ("Leave Quota Tracker", "Real-time visibility into remaining Annual (14), Medical (9), and Casual (4) leave quotas."),
            ("Direct Navigation Cards", "Clean shortcuts to submit time-off requests, view attendance logs, and inspect payslips."),
            ("Mobile-First Responsiveness", "Optimized touch targets and fluid layout tailored for field tablets and smartphones.")
        ]
    )

    # ==========================================
    # SLIDE 14: DIGITAL PAYSLIP VIEWER
    # ==========================================
    add_image_slide(
        "Digital Payslip & Compensation Viewer",
        "Staff Workflows",
        "08_employee_payslips.png",
        [
            ("Transparent Salary Breakdown", "Itemized breakdown of Basic (₹95,000), Allowances (₹30,000), and Deductions (₹12,000)."),
            ("Net Take-Home Visibility", "Displays verified monthly net take-home salary (₹1,13,000.00) disbursed via direct deposit."),
            ("Historical Archive", "Four-month historical payslip timeline (July, June, May, April) for tax and loan verification."),
            ("PDF Generation Trigger", "One-click simulated PDF download with instant toast confirmation alerts."),
            ("Encrypted Access", "Strictly isolated through session token verification to ensure confidentiality.")
        ]
    )

    # ==========================================
    # SLIDE 15: EXIM-AI COPILOT
    # ==========================================
    add_image_slide(
        "EXIM-AI Contextual Copilot",
        "AI Innovation",
        "11_exim_ai_copilot.png",
        [
            ("Role-Aware Intelligence", "Dynamically tailors responses: returns personal leave balances to Employees, and pending approvals to HR."),
            ("Natural Language Parsing", "Accurately interprets user queries regarding leaves, salary, attendance, and port hubs."),
            ("Actionable Deep-Linking", "Generates interactive CTA buttons (e.g., 'Manage Pending Leaves') for one-click navigation."),
            ("Domain Grounding", "Pre-loaded with accurate Indian hub operational data (JNPT Mumbai, Mundra Gujarat, Delhi Expressway)."),
            ("Floating In-App Widget", "Accessible across all authenticated pages with animated open/close physics.")
        ]
    )

    # ==========================================
    # SLIDE 16: COMPANY STAFF DIRECTORY
    # ==========================================
    add_image_slide(
        "Company Colleague Network Directory",
        "Team Collaboration",
        "09_company_directory.png",
        [
            ("Cross-Hub Networking", "Enables employees to look up peers, site engineers, and division heads across all Indian hubs."),
            ("Direct Contact Coordinates", "Surfaces verified work emails (@exim2d.in) and mobile phone numbers for field coordination."),
            ("Real-Time Filter Input", "Instant debounce search filtering staff by name, department, or job role."),
            ("Department Avatar Badges", "Custom color-coded avatar initials for quick visual identification."),
            ("Security Policy Compliance", "Read-only interface preventing unauthorized tampering with corporate roster records.")
        ]
    )

    # ==========================================
    # SLIDE 17: FUTURE SCOPE
    # ==========================================
    s17 = prs.slides.add_slide(blank_layout)
    add_header(s17, "Future Scope & Technological Roadmap", "System Evolution")
    add_footer(s17)

    f_items = [
        ("IoT Geofenced Punch-In", "Automated GPS-backed mobile check-in logging shifts the moment workers enter container terminal perimeters."),
        ("Direct Banking API Integration", "Direct integration with corporate banking gateways (NEFT/RTGS/SWIFT) for automated single-click salary disbursement."),
        ("LLM Compliance Policy Drafter", "Generative AI models fine-tuned to draft customized cross-border trade contracts and occupational safety policies."),
        ("Predictive Attrition & Absence Modeling", "Machine learning forecasting algorithms predicting seasonal labor shortages prior to peak trade quarters."),
        ("Biometric Hardware Terminal Sync", "Direct MQTT/WebSocket communication syncing on-site fingerprint and facial scanners at port terminals into MongoDB.")
    ]

    for i, (title, desc) in enumerate(f_items):
        box = s17.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8 + i * 0.95), Inches(11.7), Inches(0.82))
        box.fill.solid()
        box.fill.fore_color.rgb = COLOR_WHITE
        box.line.color.rgb = COLOR_CARD_BORDER
        tf = box.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = f"🚀  {title}"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEAL

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_PRIMARY
        p2.space_before = Pt(2)

    # ==========================================
    # SLIDE 18: Q & A
    # ==========================================
    s18 = prs.slides.add_slide(blank_layout)
    bg18 = s18.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg18.fill.solid()
    bg18.fill.fore_color.rgb = COLOR_PRIMARY
    bg18.line.fill.background()

    tb18 = s18.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.333), Inches(3.0))
    tf18 = tb18.text_frame
    tf18.word_wrap = True
    p = tf18.paragraphs[0]
    p.text = "QUESTIONS & ANSWERS"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE

    p2 = tf18.add_paragraph()
    p2.text = "We are now open for evaluation questions and technical discussion."
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(16)
    p2.font.color.rgb = COLOR_TEAL
    p2.space_before = Pt(16)

    p3 = tf18.add_paragraph()
    p3.text = "Department of Computer Science & Engineering\nSwami Keshvanand Institute of Technology, Management & Gramothan, Jaipur"
    p3.alignment = PP_ALIGN.CENTER
    p3.font.size = Pt(13)
    p3.font.color.rgb = RGBColor(148, 163, 184)
    p3.space_before = Pt(24)

    # ==========================================
    # SLIDE 19: THANK YOU
    # ==========================================
    s19 = prs.slides.add_slide(blank_layout)
    bg19 = s19.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg19.fill.solid()
    bg19.fill.fore_color.rgb = COLOR_PRIMARY
    bg19.line.fill.background()

    tb19 = s19.shapes.add_textbox(Inches(1.5), Inches(1.8), Inches(10.333), Inches(4.0))
    tf19 = tb19.text_frame
    tf19.word_wrap = True
    p = tf19.paragraphs[0]
    p.text = "THANK YOU!"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE

    p2 = tf19.add_paragraph()
    p2.text = "EXIM 2D World Traders HR ERP & AI Copilot Platform"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(18)
    p2.font.color.rgb = COLOR_ACCENT
    p2.space_before = Pt(14)

    p3 = tf19.add_paragraph()
    p3.text = "GitHub Repository: https://github.com/khushikmandal2005-commits/-hr-erp-"
    p3.alignment = PP_ALIGN.CENTER
    p3.font.size = Pt(14)
    p3.font.color.rgb = COLOR_TEAL
    p3.space_before = Pt(14)

    p4 = tf19.add_paragraph()
    p4.text = "Presented By: [Friend's Name] & [Your Name]\nSKIT, Jaipur  •  www.skit.ac.in"
    p4.alignment = PP_ALIGN.CENTER
    p4.font.size = Pt(13)
    p4.font.color.rgb = RGBColor(148, 163, 184)
    p4.space_before = Pt(18)

    # Save presentation
    output_path = os.path.join(base_dir, "EXIM_2D_HR_ERP_Presentation.pptx")
    prs.save(output_path)
    print(f"SUCCESS: Presentation saved at: {output_path}")

if __name__ == "__main__":
    create_deck()
