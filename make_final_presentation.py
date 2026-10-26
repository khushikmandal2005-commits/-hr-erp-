import os
import shutil
import zipfile
from pptx import Presentation
from PIL import Image

def generate_official_ppt():
    base_dir = os.path.abspath(os.path.dirname(__file__))
    template_pptx = 'C:/Users/Khushi/.gemini/antigravity/brain/f5fd6042-0789-4aed-a6bf-d55aa2c2be5d/.user_uploaded/media_1792683463122_0e578dd9.pptx'
    step1_pptx = os.path.join(base_dir, 'step1_text_updated.pptx')
    final_output = os.path.join(base_dir, 'EXIM_2D_HR_ERP_Official_Presentation.pptx')
    ss_dir = os.path.join(base_dir, 'presentation_screenshots')

    print("Step 1: Updating text in template presentation...")
    prs = Presentation(template_pptx)

    # Slide 1: Names and Topic
    s1 = prs.slides[0]
    for shape in s1.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text
            if 'Mahak' in t:
                shape.text_frame.text = "Submitted By:\nName:- Khushi Kumari\nName:- Lakshay Chanwariya"
            elif 'TOPIC:' in t:
                shape.text_frame.text = "TOPIC:  EXIM 2D HR ERP"

    # Slide 3: Company Introduction
    s3 = prs.slides[2]
    for shape in s3.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text
            if 'Tecorb' in t or 'Compucom' in t or 'leading' in t.lower():
                shape.text_frame.text = (
                    "Exim 2D World Traders Pvt. Ltd. is a leading enterprise operating in global import-export trade, maritime cargo logistics, container fleet management, and large-scale highway civil infrastructure projects.\n\n"
                    "The company operates major logistics and field project hubs across India—including JNPT Port Mumbai, Mundra Port Gujarat, Chennai Port Hub, Kolkata Logistics Center, and Corporate Headquarters in New Delhi.\n\n"
                    "Exim 2D World Traders employs over 140+ multidisciplinary personnel comprising port logistics managers, customs clearance specialists, civil site engineers, and corporate management across multiple timezones.\n\n"
                    "The software & ERP division emphasizes digital workforce modernization, quality, and technology adoption, providing a strong platform for learning, skill development, and professional growth in modern Web Development."
                )

    # Slide 6: Goal of Training
    s6 = prs.slides[5]
    for shape in s6.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text
            if 'hands-on' in t.lower() or 'tecorb' in t.lower() or 'devenue' in t.lower():
                shape.text_frame.text = (
                    "1. To gain hands-on experience with full-stack web development using modern decoupled architectures (React 19, Node.js, Express.js, MongoDB).\n"
                    "2. To learn end-to-end enterprise project development, including shift biometrics, role-based workflows, and secure APIs using JWT.\n"
                    "3. To design and implement a Dual-Persona portal separating Executive HR management and Staff Self-Service (ESS) within a single SPA.\n"
                    "4. To develop professional skills such as effective teamwork, communication, and collaboration in an industrial project environment.\n"
                    "5. To strengthen expertise in version control (Git & GitHub), API testing, and deployment workflows while applying all skills to build EXIM 2D HR ERP platform."
                )

    # Slide 7: Overview
    s7 = prs.slides[6]
    for shape in s7.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text
            if 'App Review' in t or 'Overview' in t:
                shape.text_frame.text = "EXIM 2D HR ERP Overview"
            elif 'TecOrb' in t or 'unstructured' in t or 'dual-engine' in t.lower():
                shape.text_frame.text = (
                    "The EXIM 2D HR ERP & In-App AI Copilot Platform is an enterprise-grade full-stack web platform developed for Exim 2D World Traders Pvt. Ltd. to solve the challenge of managing a distributed global workforce across international shipping ports, civil infrastructure sites, and corporate offices.\n\n"
                    "It operates on a dual-engine architecture combining an Executive HR Command Center (for real-time KPI tracking, biometric shift logs, payroll disbursements, and department management) with an Employee Self-Service (ESS) Portal (for shift clock-in, leave quotas, and digital payslips).\n\n"
                    "The entire system is augmented by EXIM-AI, an intelligent conversational copilot that parses natural language queries and automatically surfaces one-click navigation actions for leaves, attendance, and port logistics status."
                )

    # Section Headers for Slides 8 to 20
    headers_map = {
        7: "SAMPLE HR DATASET SECTION",
        8: "GATEWAY SWITCHER SECTION",
        9: "LOGIN & AUTH SECTION",
        10: "HR DASHBOARD & RADAR SECTION",
        11: "EMPLOYEES DIRECTORY SECTION",
        12: "ATTENDANCE BIOMETRICS SECTION",
        13: "ENTERPRISE PAYROLL SECTION",
        14: "EMPLOYEE SELF-SERVICE SECTION",
        15: "DIGITAL PAYSLIP SECTION",
        16: "EXIM-AI COPILOT CHATBOT SECTION",
        17: "COLLEAGUE NETWORK SECTION",
        18: "CERTIFICATE - LAKSHAY CHANWARIYA",
        19: "API TESTING & SYSTEM ARCHITECTURE"
    }

    for idx, new_title in headers_map.items():
        if idx < len(prs.slides):
            slide = prs.slides[idx]
            for shape in slide.shapes:
                if shape.has_text_frame:
                    t = shape.text_frame.text.strip()
                    if t in ['REVIEW SECTION', 'AI CHATBOT', 'UPLOAD DATASET', 'DATASET INVENTORY', 'UPLOAD AND INGEST DATASET', 'RELEASE NOTES', 'API TESTING']:
                        shape.text_frame.text = new_title

    # Slide 21: Future Scope
    s21 = prs.slides[20]
    for shape in s21.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text
            if 'Scraping' in t or 'Sentiment' in t or 'Workflow' in t:
                shape.text_frame.text = (
                    "1. Real-Time IoT Geofencing for Automated Port Punch-In.\n\n"
                    "2. Direct Corporate Banking Gateway Sync (Automated Salary Disbursal).\n\n"
                    "3. LLM Compliance & Cross-Border Employment Contract Generation.\n\n"
                    "4. Predictive Attrition & Seasonal Workforce Absence Forecasting.\n\n"
                    "5. Hardware Biometric Fingerprint & Facial Terminal Integration."
                )

    prs.save(step1_pptx)
    print("Step 1 Complete: Text updated and saved to step1_text_updated.pptx")

    # Step 2: Unzip and Replace Media Images
    print("Step 2: Replacing media images...")
    work_dir = os.path.join(base_dir, "temp_pptx_pack")
    if os.path.exists(work_dir):
        shutil.rmtree(work_dir)
    os.makedirs(work_dir, exist_ok=True)

    with zipfile.ZipFile(step1_pptx, 'r') as z:
        z.extractall(work_dir)

    media_dir = os.path.join(work_dir, 'ppt', 'media')

    image_replacements = {
        'image4.jpeg': os.path.join(ss_dir, 'cert_khushi.png'),       # Slide 4: Khushi Kumari Certificate
        'image6.jpeg': os.path.join(ss_dir, '10_sample_dataset_excel.png'), # Slide 8: Sample HR Dataset
        'image7.jpeg': os.path.join(ss_dir, '01_landing_gateway.png'),     # Slide 9: Gateway Switcher
        'image8.jpeg': os.path.join(ss_dir, '02_hr_login.png'),           # Slide 10: Login
        'image9.jpeg': os.path.join(ss_dir, '03_hr_dashboard.png'),      # Slide 11: HR Dashboard
        'image10.png': os.path.join(ss_dir, '04_employees_directory.png'),# Slide 12: Employees Directory
        'image11.jpeg': os.path.join(ss_dir, '05_attendance_sheet.png'),   # Slide 13: Attendance
        'image12.jpeg': os.path.join(ss_dir, '06_payroll_management.png'), # Slide 14: Payroll
        'image13.jpeg': os.path.join(ss_dir, '07_employee_dashboard.png'), # Slide 15: Employee Self-Service
        'image14.jpeg': os.path.join(ss_dir, '08_employee_payslips.png'),  # Slide 16: Payslips
        'image15.png': os.path.join(ss_dir, '11_exim_ai_copilot.png'),    # Slide 17: EXIM-AI Copilot
        'image16.jpeg': os.path.join(ss_dir, '09_company_directory.png'),  # Slide 18: Colleague Network
        'image17.jpeg': os.path.join(ss_dir, 'cert_lakshay.png'),          # Slide 19: Lakshay Chanwariya Certificate
        'image18.jpeg': os.path.join(ss_dir, '03_hr_dashboard.png'),      # Slide 20: Architecture / Dashboard
    }

    for target_name, src_path in image_replacements.items():
        if os.path.exists(src_path):
            target_path = os.path.join(media_dir, target_name)
            im = Image.open(src_path)
            if target_name.endswith('.jpeg') or target_name.endswith('.jpg'):
                im.convert('RGB').save(target_path, 'JPEG', quality=95)
            else:
                im.save(target_path, 'PNG')
            print(f"Replaced {target_name} with {os.path.basename(src_path)}")

    # Step 3: Re-pack into the final PPTX
    print("Step 3: Repacking into final presentation...")
    if os.path.exists(final_output):
        try:
            os.remove(final_output)
        except:
            pass

    with zipfile.ZipFile(final_output, 'w', zipfile.ZIP_DEFLATED) as z_out:
        for root, dirs, files in os.walk(work_dir):
            for file in files:
                abs_path = os.path.join(root, file)
                rel_path = os.path.relpath(abs_path, work_dir)
                z_out.write(abs_path, rel_path)

    # Clean up temp directories
    shutil.rmtree(work_dir, ignore_errors=True)
    if os.path.exists(step1_pptx):
        os.remove(step1_pptx)

    print(f"COMPLETE! Official Presentation generated at: {final_output}")

if __name__ == '__main__':
    generate_official_ppt()
