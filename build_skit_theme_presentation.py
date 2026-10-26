import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    base_dir = os.path.abspath(os.path.dirname(__file__))
    img_dir = os.path.join(base_dir, "presentation_screenshots")
    assets_dir = os.path.join(base_dir, "extracted_template_assets", "ppt", "media")

    skit_logo = os.path.join(assets_dir, "image1.png")
    swami_img = os.path.join(assets_dir, "image2.png")
    cert_khushi = os.path.join(img_dir, "cert_khushi.png")
    cert_lakshay = os.path.join(img_dir, "cert_lakshay.png")

    COLOR_RED = RGBColor(192, 0, 0)       # #C00000 - Exact SKIT Red
    COLOR_MAROON = RGBColor(128, 0, 0)    # #800000 - Title Slide Heading
    COLOR_BLACK = RGBColor(0, 0, 0)
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_GRAY = RGBColor(70, 70, 70)

    # Helper: Add Standard SKIT Header and Footer
    def add_skit_chrome(slide, title_text):
        # 1. SKIT Logo on top-left
        if os.path.exists(skit_logo):
            slide.shapes.add_picture(skit_logo, Inches(0.39), Inches(0.19), Inches(1.24), Inches(1.31))

        # 2. Top Red Banner
        banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2.47), Inches(0.24), Inches(10.30), Inches(0.68))
        banner.fill.solid()
        banner.fill.fore_color.rgb = COLOR_RED
        banner.line.fill.background()
        tf = banner.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE
        p.alignment = PP_ALIGN.CENTER

        # 3. Bottom Red Footer Bar
        footer = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.0), Inches(7.23), Inches(13.333), Inches(0.27))
        footer.fill.solid()
        footer.fill.fore_color.rgb = COLOR_RED
        footer.line.fill.background()

        # 4. Footer Left Text
        f_left = slide.shapes.add_textbox(Inches(0.15), Inches(7.20), Inches(2.5), Inches(0.3))
        tf_l = f_left.text_frame
        p_l = tf_l.paragraphs[0]
        p_l.text = "SKIT, JAIPUR"
        p_l.font.size = Pt(9)
        p_l.font.bold = True
        p_l.font.color.rgb = COLOR_WHITE

        # 5. Footer Center Text
        f_mid = slide.shapes.add_textbox(Inches(5.5), Inches(7.20), Inches(3.0), Inches(0.3))
        tf_m = f_mid.text_frame
        p_m = tf_m.paragraphs[0]
        p_m.text = "www.skit.ac.in"
        p_m.font.size = Pt(9)
        p_m.font.bold = True
        p_m.font.color.rgb = COLOR_WHITE
        p_m.alignment = PP_ALIGN.CENTER

    # Helper: Add Section Image Slide (like Page 8, 9, 11 of sample)
    def add_section_image_slide(title_text, img_filename):
        s = prs.slides.add_slide(blank_layout)
        # Red Header & Footer
        add_skit_chrome(s, title_text)

        img_path = os.path.join(img_dir, img_filename)
        if os.path.exists(img_path):
            # Centered screenshot image
            s.shapes.add_picture(img_path, Inches(1.5), Inches(1.2), Inches(10.33), Inches(5.8))

    # =========================================================================
    # SLIDE 1: TITLE SLIDE (Exact replica of sample Page 1)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    # SKIT Logo top-left
    if os.path.exists(skit_logo):
        s1.shapes.add_picture(skit_logo, Inches(0.39), Inches(0.19), Inches(1.24), Inches(1.31))
    # Swami Keshvanand Ji picture top-right
    if os.path.exists(swami_img):
        s1.shapes.add_picture(swami_img, Inches(12.0), Inches(0.19), Inches(0.85), Inches(1.38))

    # College Name Header (Maroon, Centered)
    c_header = s1.shapes.add_textbox(Inches(1.8), Inches(0.35), Inches(10.0), Inches(1.1))
    tf = c_header.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = "Swami Keshvanand Institute of Technology, Management & Gramothan, Jaipur"
    p1.font.size = Pt(16)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_MAROON
    p1.alignment = PP_ALIGN.CENTER

    p2 = tf.add_paragraph()
    p2.text = "(An Autonomous Institute Affiliated to Rajasthan Technical University, Kota)"
    p2.font.size = Pt(13)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_MAROON
    p2.alignment = PP_ALIGN.CENTER

    # Middle Title
    m_box = s1.shapes.add_textbox(Inches(2.0), Inches(2.2), Inches(9.33), Inches(2.2))
    tf_m = m_box.text_frame
    tf_m.word_wrap = True
    pm1 = tf_m.paragraphs[0]
    pm1.text = "Industrial Training\nPresentation"
    pm1.font.size = Pt(22)
    pm1.font.bold = True
    pm1.font.color.rgb = COLOR_BLACK
    pm1.alignment = PP_ALIGN.CENTER

    pm2 = tf_m.add_paragraph()
    pm2.text = "TOPIC:  EXIM 2D HR ERP"
    pm2.font.size = Pt(32)
    pm2.font.bold = True
    pm2.font.underline = True
    pm2.font.color.rgb = COLOR_BLACK
    pm2.alignment = PP_ALIGN.CENTER
    pm2.space_before = Pt(18)

    # Bottom Left: Submitted To
    sub_to = s1.shapes.add_textbox(Inches(0.47), Inches(5.0), Inches(4.5), Inches(1.8))
    tf_to = sub_to.text_frame
    tf_to.word_wrap = True
    pt1 = tf_to.paragraphs[0]
    pt1.text = "Submitted To:"
    pt1.font.size = Pt(15)
    pt1.font.bold = True
    pt1.font.color.rgb = COLOR_BLACK

    pt2 = tf_to.add_paragraph()
    pt2.text = "Name:- Dr. Meenakshi Nawal"
    pt2.font.size = Pt(14)
    pt2.font.bold = True
    pt2.font.color.rgb = COLOR_BLACK
    pt2.space_before = Pt(6)

    pt3 = tf_to.add_paragraph()
    pt3.text = "Associate Professor\nDepartment of Computer Science & Engineering"
    pt3.font.size = Pt(13)
    pt3.font.color.rgb = COLOR_BLACK

    # Bottom Right: Submitted By (Khushi Kumari & Lakshay Chanwariya)
    sub_by = s1.shapes.add_textbox(Inches(8.5), Inches(5.0), Inches(4.5), Inches(1.8))
    tf_by = sub_by.text_frame
    tf_by.word_wrap = True
    pb1 = tf_by.paragraphs[0]
    pb1.text = "Submitted By:"
    pb1.font.size = Pt(15)
    pb1.font.bold = True
    pb1.font.color.rgb = COLOR_BLACK

    pb2 = tf_by.add_paragraph()
    pb2.text = "Name:- Khushi Kumari\nName:- Lakshay Chanwariya"
    pb2.font.size = Pt(14)
    pb2.font.bold = True
    pb2.font.color.rgb = COLOR_BLACK
    pb2.space_before = Pt(6)

    pb3 = tf_by.add_paragraph()
    pb3.text = "Department of Computer Science & Engineering\nSKIT, Jaipur"
    pb3.font.size = Pt(12)
    pb3.font.color.rgb = COLOR_BLACK

    # =========================================================================
    # SLIDE 2: INDEX
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_skit_chrome(s2, "Index")

    idx_box = s2.shapes.add_textbox(Inches(2.4), Inches(1.3), Inches(8.5), Inches(5.6))
    tf_i = idx_box.text_frame
    tf_i.word_wrap = True

    index_items = [
        "1.  Company Introduction",
        "2.  Certificate",
        "3.  Technology Used",
        "4.  Goal of Training",
        "5.  Project Overview",
        "6.  Working Project",
        "7.  Future Scope",
        "8.  Q&A"
    ]
    for i, itm in enumerate(index_items):
        p = tf_i.paragraphs[0] if i == 0 else tf_i.add_paragraph()
        p.text = itm
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = COLOR_BLACK
        if i > 0:
            p.space_before = Pt(16)

    # =========================================================================
    # SLIDE 3: COMPANY INTRODUCTION (EXIM 2D WORLD TRADERS PVT. LTD.)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_skit_chrome(s3, "Company Introduction")

    c_box = s3.shapes.add_textbox(Inches(2.2), Inches(1.3), Inches(9.8), Inches(5.6))
    tf_c = c_box.text_frame
    tf_c.word_wrap = True

    c_paragraphs = [
        ("Exim 2D World Traders Pvt. Ltd.", " is a prominent international enterprise specializing in global import-export trade, maritime cargo logistics, container fleet management, and large-scale civil infrastructure projects."),
        ("Corporate Operations: ", "The company operates major logistics and field project hubs across India—including JNPT Port Mumbai, Mundra Port Gujarat, Chennai Port Hub, Kolkata Logistics Center, and Corporate Headquarters in New Delhi."),
        ("Multidisciplinary Workforce: ", "Exim 2D World Traders employs over 140+ multidisciplinary personnel comprising port logistics managers, customs clearance specialists, civil site engineers, and corporate management across timezones."),
        ("Software & ERP Division: ", "The company focuses on technology adoption, digital workforce modernization, and customer satisfaction, providing a rigorous and supportive platform for learning, skill development, and professional growth in modern Web Development.")
    ]
    for i, (bold_txt, normal_txt) in enumerate(c_paragraphs):
        p = tf_c.paragraphs[0] if i == 0 else tf_c.add_paragraph()
        p.text = bold_txt
        p.font.bold = True
        p.font.size = Pt(16)
        p.font.color.rgb = COLOR_BLACK
        if i > 0: p.space_before = Pt(18)

        run = p.add_run()
        run.text = normal_txt
        run.font.bold = False
        run.font.size = Pt(16)
        run.font.color.rgb = COLOR_BLACK

    # =========================================================================
    # SLIDE 4: CERTIFICATE - KHUSHI KUMARI
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_skit_chrome(s4, "Certificate - Khushi Kumari")
    if os.path.exists(cert_khushi):
        s4.shapes.add_picture(cert_khushi, Inches(2.8), Inches(1.2), Inches(7.73), Inches(5.75))

    # =========================================================================
    # SLIDE 5: CERTIFICATE - LAKSHAY CHANWARIYA
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_skit_chrome(s5, "Certificate - Lakshay Chanwariya")
    if os.path.exists(cert_lakshay):
        s5.shapes.add_picture(cert_lakshay, Inches(2.8), Inches(1.2), Inches(7.73), Inches(5.75))

    # =========================================================================
    # SLIDE 6: TECHNOLOGY STACK USED
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_skit_chrome(s6, "Technology Stack Used")

    t_box = s6.shapes.add_textbox(Inches(1.8), Inches(1.4), Inches(10.0), Inches(5.5))
    tf_t = t_box.text_frame
    tf_t.word_wrap = True

    tech_categories = [
        ("Frontend Architecture:", " React 19, Vite 8, Tailwind CSS v4, Framer Motion 12, Recharts 3, React Router v7."),
        ("Backend Services:", " Node.js, Express.js RESTful API framework, Axios HTTP Client."),
        ("Database & Ingestion:", " MongoDB (NoSQL), Mongoose ODM, CSV/JSON Seeding Pipeline (seedData.js)."),
        ("Security & Authentication:", " Stateless JWT (JSON Web Tokens), bcrypt hashing, Role-Based Access Control (RBAC)."),
        ("AI & Intelligent Copilot:", " EXIM-AI Contextual Engine, Natural Language Intent Parser, Deep-Link Action Dispatcher.")
    ]
    for i, (cat, val) in enumerate(tech_categories):
        p = tf_t.paragraphs[0] if i == 0 else tf_t.add_paragraph()
        p.text = f"•  {cat}"
        p.font.bold = True
        p.font.size = Pt(17)
        p.font.color.rgb = COLOR_RED
        if i > 0: p.space_before = Pt(18)

        run = p.add_run()
        run.text = val
        run.font.bold = False
        run.font.size = Pt(16)
        run.font.color.rgb = COLOR_BLACK

    # =========================================================================
    # SLIDE 7: GOAL OF TRAINING
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_skit_chrome(s7, "Goal Of Training")

    g_box = s7.shapes.add_textbox(Inches(1.8), Inches(1.2), Inches(10.2), Inches(5.8))
    tf_g = g_box.text_frame
    tf_g.word_wrap = True

    goals = [
        ("1. To gain hands-on experience with full-stack web development", " using modern decoupled architectures (React 19, Node.js, Express.js, MongoDB)."),
        ("2. To learn end-to-end enterprise project development,", " including shift biometrics, role-based workflows, and secure APIs using JWT and RBAC."),
        ("3. To design and implement a Dual-Persona portal", " separating Executive HR management and Staff Self-Service (ESS) within a single SPA."),
        ("4. To develop professional skills", " such as effective teamwork, communication, and collaboration in an industrial project environment."),
        ("5. To strengthen expertise in version control (Git & GitHub), API testing, and deployment workflows", " while applying all skills to build EXIM 2D HR ERP platform.")
    ]
    for i, (bold_txt, normal_txt) in enumerate(goals):
        p = tf_g.paragraphs[0] if i == 0 else tf_g.add_paragraph()
        p.text = bold_txt
        p.font.bold = True
        p.font.size = Pt(15)
        p.font.color.rgb = COLOR_BLACK
        if i > 0: p.space_before = Pt(14)

        run = p.add_run()
        run.text = normal_txt
        run.font.bold = False
        run.font.size = Pt(14)
        run.font.color.rgb = COLOR_BLACK

    # =========================================================================
    # SLIDE 8: PROJECT OVERVIEW
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_skit_chrome(s8, "EXIM 2D HR ERP Overview")

    o_box = s8.shapes.add_textbox(Inches(1.8), Inches(1.3), Inches(10.2), Inches(5.6))
    tf_o = o_box.text_frame
    tf_o.word_wrap = True

    overview_points = [
        ("The EXIM 2D HR ERP Platform", " is an enterprise-grade full-stack web application developed for Exim 2D World Traders Pvt. Ltd. to streamline workforce governance across international shipping ports, civil infrastructure sites, and corporate offices."),
        ("Dual-Persona Portal Architecture: ", "It operates on a decoupled architecture combining an Executive HR Command Center (for real-time KPI monitoring, biometric shift logs, payroll disbursements, and department management) with an Employee Self-Service (ESS) Portal (for shift clock-in, leave balances, and digital payslips)."),
        ("Contextual AI Copilot (EXIM-AI): ", "The platform is integrated with EXIM-AI, an intelligent conversational copilot that interprets natural language staff queries and provides automated one-click deep-link navigation actions for leaves, attendance, and port operations."),
        ("Enterprise Indian Workforce Grounding: ", "The system is populated with realistic Indian workforce data across key maritime and trade hubs—including JNPT Port Mumbai, Mundra Port Gujarat, Chennai Port, Kolkata Logistics Hub, and Delhi Headquarters.")
    ]
    for i, (bold_txt, normal_txt) in enumerate(overview_points):
        p = tf_o.paragraphs[0] if i == 0 else tf_o.add_paragraph()
        p.text = bold_txt
        p.font.bold = True
        p.font.size = Pt(15)
        p.font.color.rgb = COLOR_BLACK
        if i > 0: p.space_before = Pt(16)

        run = p.add_run()
        run.text = normal_txt
        run.font.bold = False
        run.font.size = Pt(15)
        run.font.color.rgb = COLOR_BLACK

    # =========================================================================
    # WORKING PROJECT SLIDES (UI SCREENSHOTS - SAMPLE PPT THEME)
    # =========================================================================
    add_section_image_slide("SAMPLE HR DATASET SECTION", "10_sample_dataset_excel.png")
    add_section_image_slide("GATEWAY SWITCHER SECTION", "01_landing_gateway.png")
    add_section_image_slide("LOGIN SECTION", "02_hr_login.png")
    add_section_image_slide("HR DASHBOARD & RADAR SECTION", "03_hr_dashboard.png")
    add_section_image_slide("EMPLOYEES DIRECTORY SECTION", "04_employees_directory.png")
    add_section_image_slide("ATTENDANCE BIOMETRICS SECTION", "05_attendance_sheet.png")
    add_section_image_slide("ENTERPRISE PAYROLL SECTION", "06_payroll_management.png")
    add_section_image_slide("EMPLOYEE SELF-SERVICE SECTION", "07_employee_dashboard.png")
    add_section_image_slide("DIGITAL PAYSLIP SECTION", "08_employee_payslips.png")
    add_section_image_slide("EXIM-AI COPILOT CHATBOT SECTION", "11_exim_ai_copilot.png")
    add_section_image_slide("COLLEAGUE NETWORK SECTION", "09_company_directory.png")

    # =========================================================================
    # SLIDE 20: FUTURE SCOPE
    # =========================================================================
    s20 = prs.slides.add_slide(blank_layout)
    add_skit_chrome(s20, "Future Scope")

    fs_box = s20.shapes.add_textbox(Inches(1.8), Inches(1.3), Inches(10.2), Inches(5.6))
    tf_fs = fs_box.text_frame
    tf_fs.word_wrap = True

    fs_items = [
        "1.  Real-Time IoT Geofencing for Automated Port Punch-In.",
        "2.  Direct Corporate Banking Gateway Integration (Automated Salary Disbursal).",
        "3.  LLM Compliance & Cross-Border Employment Contract Generation.",
        "4.  Predictive Attrition & Seasonal Workforce Absence Forecasting.",
        "5.  Hardware Biometric Fingerprint & Facial Terminal Integration."
    ]
    for i, itm in enumerate(fs_items):
        p = tf_fs.paragraphs[0] if i == 0 else tf_fs.add_paragraph()
        p.text = itm
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = COLOR_BLACK
        if i > 0: p.space_before = Pt(20)

    # =========================================================================
    # SLIDE 21: Q & A (Exact match of sample Page 23)
    # =========================================================================
    s21 = prs.slides.add_slide(blank_layout)
    # SKIT Logo top-left
    if os.path.exists(skit_logo):
        s21.shapes.add_picture(skit_logo, Inches(0.39), Inches(0.19), Inches(1.24), Inches(1.31))

    # Bottom Red Footer Bar
    footer21 = s21.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.0), Inches(7.23), Inches(13.333), Inches(0.27))
    footer21.fill.solid()
    footer21.fill.fore_color.rgb = COLOR_RED
    footer21.line.fill.background()

    fl21 = s21.shapes.add_textbox(Inches(0.15), Inches(7.20), Inches(2.5), Inches(0.3))
    fl21.text_frame.paragraphs[0].text = "SKIT, JAIPUR"
    fl21.text_frame.paragraphs[0].font.size = Pt(9)
    fl21.text_frame.paragraphs[0].font.bold = True
    fl21.text_frame.paragraphs[0].font.color.rgb = COLOR_WHITE

    fm21 = s21.shapes.add_textbox(Inches(5.5), Inches(7.20), Inches(3.0), Inches(0.3))
    fm21.text_frame.paragraphs[0].text = "www.skit.ac.in"
    fm21.text_frame.paragraphs[0].font.size = Pt(9)
    fm21.text_frame.paragraphs[0].font.bold = True
    fm21.text_frame.paragraphs[0].font.color.rgb = COLOR_WHITE
    fm21.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    # Centered Red Q & A text
    qa_box = s21.shapes.add_textbox(Inches(2.0), Inches(2.6), Inches(9.33), Inches(2.0))
    p_qa = qa_box.text_frame.paragraphs[0]
    p_qa.text = "Q & A"
    p_qa.font.size = Pt(64)
    p_qa.font.bold = True
    p_qa.font.color.rgb = COLOR_RED
    p_qa.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 22: THANK YOU (Exact match of sample Page 24)
    # =========================================================================
    s22 = prs.slides.add_slide(blank_layout)
    # SKIT Logo top-left
    if os.path.exists(skit_logo):
        s22.shapes.add_picture(skit_logo, Inches(0.39), Inches(0.19), Inches(1.24), Inches(1.31))

    # Bottom Red Footer Bar
    footer22 = s22.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.0), Inches(7.23), Inches(13.333), Inches(0.27))
    footer22.fill.solid()
    footer22.fill.fore_color.rgb = COLOR_RED
    footer22.line.fill.background()

    fl22 = s22.shapes.add_textbox(Inches(0.15), Inches(7.20), Inches(2.5), Inches(0.3))
    fl22.text_frame.paragraphs[0].text = "SKIT, JAIPUR"
    fl22.text_frame.paragraphs[0].font.size = Pt(9)
    fl22.text_frame.paragraphs[0].font.bold = True
    fl22.text_frame.paragraphs[0].font.color.rgb = COLOR_WHITE

    fm22 = s22.shapes.add_textbox(Inches(5.5), Inches(7.20), Inches(3.0), Inches(0.3))
    fm22.text_frame.paragraphs[0].text = "www.skit.ac.in"
    fm22.text_frame.paragraphs[0].font.size = Pt(9)
    fm22.text_frame.paragraphs[0].font.bold = True
    fm22.text_frame.paragraphs[0].font.color.rgb = COLOR_WHITE
    fm22.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    # Centered Red Thank You text
    ty_box = s22.shapes.add_textbox(Inches(2.0), Inches(2.4), Inches(9.33), Inches(2.2))
    p_ty = ty_box.text_frame.paragraphs[0]
    p_ty.text = "Thank\nYou"
    p_ty.font.size = Pt(56)
    p_ty.font.bold = True
    p_ty.font.color.rgb = COLOR_RED
    p_ty.alignment = PP_ALIGN.CENTER

    output_path = os.path.join(base_dir, "EXIM_2D_HR_ERP_SKIT_Presentation.pptx")
    prs.save(output_path)
    print(f"SUCCESS: SKIT-Themed Presentation generated at: {output_path}")

    # Also try to save to the original name if not locked
    try:
        orig_path = os.path.join(base_dir, "EXIM_2D_HR_ERP_Presentation.pptx")
        prs.save(orig_path)
        print(f"SUCCESS: Also updated original file: {orig_path}")
    except Exception as e:
        print(f"Original file was open/locked, but new file saved at: {output_path}")

if __name__ == "__main__":
    build_presentation()
